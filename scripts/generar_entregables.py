"""Llena la plantilla EN BLANCO del diccionario de datos con el catálogo, el diccionario y la bitácora.
Ejemplos y rangos se calculan desde el lago real, así el Excel no puede contradecir a los datos.

Uso: python scripts/generar_entregables.py [ruta_plantilla] [ruta_salida]
"""
import sys, os, copy, json, re
import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
from comun import RAIZ, LAGO
from catalogo import FUENTES
from diccionario import DIC, META, EXCLUIDOS, PREGUNTAS
from verificar_lago import campos_del_lago
from bitacora import BITACORA, GRUPO

PADRE = os.path.dirname(RAIZ)
PLANTILLA = sys.argv[1] if len(sys.argv) > 1 else os.path.join(PADRE, 'Diccionario_de_datos_TallerDatos.xlsx')
SALIDA = sys.argv[2] if len(sys.argv) > 2 else os.path.join(PADRE, 'Diccionario_de_datos_TallerDatos_NuevaYork.xlsx')
ULTIMA = 300   # la plantilla trae 40 filas (5–44); el lago tiene más campos, se amplían filas, listas y fórmulas


def ampliar(ws, ultima_plantilla=44):
    modelo = [copy.copy(ws.cell(ultima_plantilla, c)._style) for c in range(1, ws.max_column + 1)]
    for r in range(ultima_plantilla + 1, ULTIMA + 1):
        for c, st in enumerate(modelo, 1):
            ws.cell(r, c)._style = copy.copy(st)
    for dv in ws.data_validations.dataValidation:
        dv.sqref = openpyxl.worksheet.cell_range.MultiCellRange(re.sub(r'(\D)44\b', rf'\g<1>{ULTIMA}', str(dv.sqref)))


def resumen_formulas(ws):
    for fila in ws.iter_rows():
        for c in fila:
            if isinstance(c.value, str) and c.value.startswith('='):
                c.value = re.sub(r'(\$?[A-Z]\$?)44\b', rf'\g<1>{ULTIMA}', c.value)


def fmt_num(v):
    if isinstance(v, bool):
        return 'Sí' if v else 'No'
    if isinstance(v, int):
        return f'{v:,}'.replace(',', '.')
    if isinstance(v, float):
        return f'{v:,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.').rstrip('0').rstrip(',')
    return str(v)


def ejemplo_y_rango(vals, tipo):
    vals = [v for v in vals if v is not None]
    if not vals:
        return '(sin valores en esta corrida)', '—'
    if tipo == 'Coordenada / geometría':
        return vals[0], 'Polygon / MultiPolygon / Point en WGS84'
    nums = [v for v in vals if isinstance(v, (int, float)) and not isinstance(v, bool)]
    ej = vals[len(vals) // 2] if len(vals) > 2 else vals[0]
    if nums and len(nums) == len(vals):
        return fmt_num(ej), f'{fmt_num(min(nums))} a {fmt_num(max(nums))}'
    distintos = sorted({str(v) for v in vals})
    if len(distintos) <= 7:
        return str(ej)[:120], ' / '.join(distintos)[:250]
    return str(ej)[:120], f'texto ({len(distintos)} valores distintos)'


def main():
    wb = openpyxl.load_workbook(PLANTILLA)
    # ---------------- Instrucciones (se escribe después de los dos puntos, en la misma celda amarilla)
    ins = wb['Instrucciones']
    ins['B5'] = f"Ciudad elegida: {GRUPO['ciudad']}"
    ins['B6'] = f"Integrantes: {GRUPO['integrantes']}"
    ins['B7'] = f"Enlace al repositorio o a la carpeta del lago: {GRUPO['repo']}"

    # ---------------- 1. Fuentes
    ws = wb['1. Fuentes']; ampliar(ws)
    cols = ['id', 'nombre', 'entidad', 'tipo_publicador', 'url', 'formato', 'como_se_obtiene', 'licencia', 'personas', 'cobertura', 'periodo', 'frecuencia', 'fecha_prueba', 'estado', 'como_la_encontraron', 'observaciones']
    for i, f in enumerate(FUENTES):
        for j, k in enumerate(cols, 1):
            v = f[k]
            if k == 'observaciones' and f.get('lago'):
                v = f'{v} → Alimenta: {f["lago"]}'
            ws.cell(5 + i, j, v)

    # ---------------- 2. Diccionario
    ws = wb['2. Diccionario']; ampliar(ws)
    valores = {}
    for archivo, ruta, vals in campos_del_lago():
        key = ('_meta', ruta) if ruta.startswith('_meta.') else (archivo, ruta)
        valores.setdefault(key, []).extend(vals)
    filas = []
    orden_arch = ['demografia.json', 'distritos.json', 'calidad_de_vida_311.json', 'seguridad_delitos.json', 'seguridad_vial.json', 'vivienda.json',
                  'economia_rodajes.json', 'ambiente.json', 'movilidad.json', 'escucha_interes.json', 'catalogo.json', 'verificaciones.json']
    claves = sorted(DIC, key=lambda k: (orden_arch.index(k[0]) if k[0] in orden_arch else 50, k[0]))
    for archivo, ruta in claves:
        if (archivo, ruta) not in valores:
            continue   # documentado pero no presente en esta corrida (p. ej. error de una verificación)
        fid, orig, sig, tipo, unidad, vacio, clas, transf, preg = DIC[(archivo, ruta)]
        ej, rango = ejemplo_y_rango(valores[(archivo, ruta)], tipo)
        filas.append([fid, 'lago/' + archivo, ruta, orig, sig, tipo, unidad, ej, vacio, rango, clas, transf, PREGUNTAS[preg]])
    for ruta, (fid, orig, sig, tipo, unidad, vacio, clas, transf, preg) in META.items():
        ej, rango = ejemplo_y_rango(valores.get(('_meta', ruta), []), tipo)
        filas.append(['todas', 'lago/*.json (bloque _meta)', ruta, orig, sig, tipo, unidad, ej, vacio, rango, clas, transf, PREGUNTAS[preg]])
    for fid, archivo, campo, orig, sig, tipo, unidad, vacio, clas, transf, preg in EXCLUIDOS:
        filas.append([fid, archivo, campo, orig, sig, tipo, unidad, '(no publicado)', vacio, 'texto libre', clas, transf, PREGUNTAS[preg]])
    if len(filas) > ULTIMA - 4:
        raise SystemExit(f'El diccionario tiene {len(filas)} filas; amplía ULTIMA')
    for i, fila in enumerate(filas):
        for j, v in enumerate(fila, 1):
            ws.cell(5 + i, j, v)

    # ---------------- 3. Bitácora IA
    ws = wb['3. Bitácora IA']; ampliar(ws)
    for i, b in enumerate(BITACORA):
        for j, v in enumerate([i + 2, b['herramienta'], b['encargo'], b['respuesta'], b['comprobacion'], b['resultado'], b['fuente'], b['aprendizaje']], 1):
            ws.cell(5 + i, j, v)

    resumen_formulas(wb['Resumen'])
    wb.save(SALIDA)
    print(f'✓ {SALIDA}\n  fuentes: {len(FUENTES)} · campos: {len(filas)} · bitácora: {len(BITACORA)}')


if __name__ == '__main__':
    main()
